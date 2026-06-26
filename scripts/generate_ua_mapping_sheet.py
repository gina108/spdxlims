"""Generate urinalysis semiquant mapping sheet for client review."""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "Urinalysis Mapping"

# --- styles ---
header_font   = Font(bold=True, color="FFFFFF", size=11)
header_fill   = PatternFill("solid", fgColor="2F5496")
section_font  = Font(bold=True, size=10)
section_fill  = PatternFill("solid", fgColor="D9E1F2")
passthru_fill = PatternFill("solid", fgColor="F2F2F2")
blank_fill    = PatternFill("solid", fgColor="FFF2CC")
center        = Alignment(horizontal="center", vertical="center")
left          = Alignment(horizontal="left",  vertical="center")
thin          = Side(style="thin", color="BFBFBF")
border        = Border(left=thin, right=thin, top=thin, bottom=thin)

HEADERS = [
    "Test Code", "Test Name", "Machine Output (prefix)",
    "Report Display", "Numeric Value", "Units", "Notes"
]

ROWS = [
    # (test_code, test_name, machine_prefix, display, numeric, units, notes, row_type)
    # row_type: "section" | "data" | "blank" | "passthru"
    ("LEU", "Leukocytes",      "-neg",        "Negativo",    "",     "",        "",                        "data"),
    ("LEU", "",                "+-{n}",        "Trazas",      "15",   "Leu/uL",  "",                        "data"),
    ("LEU", "",                "1+",           "Positivo+",   "70",   "Leu/uL",  "",                        "data"),
    ("LEU", "",                "2+",           "Positivo++",  "125",  "Leu/uL",  "",                        "data"),
    ("LEU", "",                "3+",           "Positivo+++", "500",  "Leu/uL",  "",                        "data"),
    ("LEU", "",                "(other)",       "",            "",     "",        "Blank — tech enters manually", "blank"),

    ("NIT", "Nitrite",         "-neg",        "Negativo",    "",     "",        "",                        "data"),
    ("NIT", "",                "+pos",         "Positivo",    "",     "",        "",                        "data"),
    ("NIT", "",                "(other)",       "",            "",     "",        "Blank — tech enters manually", "blank"),

    ("URO", "Urobilinogen",    "-0.2",        "0.2",         "0.2",  "mg/dL",   "Only known output",       "data"),

    ("PRO", "Protein",         "-neg",        "Negativo",    "",     "",        "",                        "data"),
    ("PRO", "",                "+-{n}",        "15",          "15",   "mg/dL",   "",                        "data"),
    ("PRO", "",                "1+",           "30",          "30",   "mg/dL",   "",                        "data"),
    ("PRO", "",                "2+",           "100",         "100",  "mg/dL",   "",                        "data"),
    ("PRO", "",                "3+",           "300",         "300",  "mg/dL",   "",                        "data"),
    ("PRO", "",                "(other)",       "",            "",     "",        "Blank — tech enters manually", "blank"),

    ("PH",  "Urine pH",        "(numeric)",    "(as printed)", "",     "",        "Passes through unchanged", "passthru"),

    ("BLO", "Blood",           "-neg",        "Negativo",    "",     "",        "",                        "data"),
    ("BLO", "",                "+-{n}",        "Trazas",      "10",   "Ery/uL",  "",                        "data"),
    ("BLO", "",                "1+",           "Positivo+",   "25",   "Ery/uL",  "",                        "data"),
    ("BLO", "",                "2+",           "Positivo++",  "80",   "Ery/uL",  "",                        "data"),
    ("BLO", "",                "3+",           "Positivo+++", "200",  "Ery/uL",  "",                        "data"),
    ("BLO", "",                "(other)",       "",            "",     "",        "Blank — tech enters manually", "blank"),

    ("SG",  "Specific Gravity","(numeric)",    "(as printed)", "",     "",        "Passes through unchanged", "passthru"),

    ("KET", "Ketones",         "-neg",        "Negativo",    "",     "",        "",                        "data"),
    ("KET", "",                "+-{n}",        "Trazas",      "5",    "mg/dL",   "",                        "data"),
    ("KET", "",                "1+",           "Positivo+",   "15",   "mg/dL",   "",                        "data"),
    ("KET", "",                "(other)",       "",            "",     "",        "Blank — tech enters manually", "blank"),

    ("BIL", "Bilirubin",       "-neg",        "Negativo",    "",     "",        "",                        "data"),
    ("BIL", "",                "1+",           "Positivo+",   "1",    "mg/dL",   "",                        "data"),
    ("BIL", "",                "2+",           "Positivo++",  "2",    "mg/dL",   "",                        "data"),
    ("BIL", "",                "(other)",       "",            "",     "",        "Blank — tech enters manually", "blank"),

    ("GLU", "Glucose",         "-neg",        "Negativo",    "",     "",        "",                        "data"),
    ("GLU", "",                "+-{n}",        "100",         "100",  "mg/dL",   "",                        "data"),
    ("GLU", "",                "1+",           "250",         "250",  "mg/dL",   "",                        "data"),
    ("GLU", "",                "2+",           "500",         "500",  "mg/dL",   "",                        "data"),
    ("GLU", "",                "3+",           "1000",        "1000", "mg/dL",   "",                        "data"),
    ("GLU", "",                "(other)",       "",            "",     "",        "Blank — tech enters manually", "blank"),
]

# --- write header ---
ws.row_dimensions[1].height = 22
for col, h in enumerate(HEADERS, start=1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center
    cell.border = border

# --- write data rows ---
current_test = None
row_num = 2
for entry in ROWS:
    test_code, test_name, machine, display, numeric, units, notes, row_type = entry

    # section divider when test changes
    if test_code != current_test:
        current_test = test_code
        ws.row_dimensions[row_num].height = 4
        for col in range(1, len(HEADERS) + 1):
            c = ws.cell(row=row_num, column=col, value="")
            c.fill = PatternFill("solid", fgColor="2F5496")
        row_num += 1

    ws.row_dimensions[row_num].height = 18
    values = [test_code, test_name, machine, display, numeric, units, notes]
    for col, val in enumerate(values, start=1):
        cell = ws.cell(row=row_num, column=col, value=val)
        cell.border = border
        cell.alignment = left

        if row_type == "passthru":
            cell.fill = passthru_fill
            cell.font = Font(italic=True, color="595959", size=10)
        elif row_type == "blank":
            cell.fill = blank_fill
            cell.font = Font(italic=True, color="7F6000", size=10)
        else:
            cell.font = Font(size=10)

    row_num += 1

# --- column widths ---
col_widths = [12, 18, 26, 18, 16, 10, 34]
for i, w in enumerate(col_widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

# --- legend ---
row_num += 1
ws.cell(row=row_num, column=1, value="Legend").font = Font(bold=True, size=10)
row_num += 1

legend = [
    (PatternFill("solid", fgColor="FFFFFF"), "Normal mapping — machine prefix → report label + numeric value"),
    (passthru_fill,                           "Pass-through — value printed as-is from the machine"),
    (blank_fill,                              "Unrecognised output — result left blank for tech to enter manually"),
]
for fill, text in legend:
    c1 = ws.cell(row=row_num, column=1, value="")
    c1.fill = fill
    c1.border = border
    c2 = ws.cell(row=row_num, column=2, value=text)
    c2.font = Font(size=10)
    c2.alignment = left
    row_num += 1

# --- freeze header ---
ws.freeze_panes = "A2"

out = r"C:\Users\Admin\Documents\GitHub\spdxlims\urinalysis_mapping_review.xlsx"
wb.save(out)
print(f"Saved: {out}")
